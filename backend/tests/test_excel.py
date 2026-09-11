from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from app.core.errors import ApiError
from app.domain.expenses import calculate_expense_totals
from app.excel.template_contract import EXCEL_TEMPLATE, apply_output_page_setup
from app.schemas.excel import ExcelGenerateRequest
from app.services.excel_generator import (
    ResolvedProject,
    WorkbookResult,
    build_download_filename,
    content_disposition,
    generate_expense_workbook,
)
from app.services.subsidy_calculation import (
    calculate_trip_subsidies,
    merge_overlapping_subsidy_trips,
    request_trips,
)

TEMPLATE_PATH = Path(__file__).parents[1] / "app" / "templates" / "expense_template.xlsx"
CANONICAL_TEMPLATE_SHA256 = "39ba8327577e3ea109aa462b82c0fde265f6e7fa6c50fc2a824c6890326b27c5"


def trip() -> dict[str, object]:
    return {
        "tripType": "business",
        "startDate": "2026-06-30",
        "startTime": "09:00",
        "endDate": "2026-07-07",
        "endTime": "18:00",
    }


def item(
    *,
    category: str = "other",
    date: str = "2026-06-30",
    display_date: str | None = None,
    description: str = "测试费用",
    amount: str = "1.00",
    receipt_count: int = 1,
) -> dict[str, object]:
    return {
        "category": category,
        "date": date,
        "displayDate": display_date or date,
        "description": description,
        "amount": amount,
        "receiptCount": receipt_count,
    }


def body(*, items: list[dict[str, object]] | None = None) -> dict[str, object]:
    return {
        "project": {"mode": "manual", "text": "P-001 示例项目"},
        "trip": trip(),
        "items": items or [],
    }


def generate(
    client,
    payload: dict[str, object],
    *,
    template_path: Path = TEMPLATE_PATH,
    employee_name: str = "开发测试用户",
    department_name: str = "测试部门",
) -> WorkbookResult:
    request = ExcelGenerateRequest.model_validate(payload)
    trips = request_trips(trip=request.trip, trips=request.trips)
    with client.app.state.database_session_factory() as database:
        subsidies = calculate_trip_subsidies(database, trips)
    totals = calculate_expense_totals(request.items, subsidies)
    return generate_expense_workbook(
        template_path=template_path,
        employee_name=employee_name,
        department_name=department_name,
        project=ResolvedProject(request.project.text, request.project.text),
        trip=request.trip,
        items=request.items,
        subsidy=subsidies[0] if request.trip is not None and len(subsidies) == 1 else None,
        totals=totals,
        trips=merge_overlapping_subsidy_trips(trips),
        subsidies=subsidies,
    )


def open_result(result: WorkbookResult):
    return load_workbook(BytesIO(result.content), data_only=False, keep_links=True)


def assert_template_invalid(client_factory, template_path: Path) -> None:
    client = client_factory()
    with pytest.raises(ApiError) as error:
        generate(
            client,
            {"project": {"mode": "manual", "text": "项目"}, "items": []},
            template_path=template_path,
        )
    assert (error.value.code, error.value.message, error.value.status_code) == (
        "EXCEL_TEMPLATE_INVALID",
        "Excel 模板结构无效，请联系管理员",
        500,
    )


def assert_archive_has_no_active_content(content: bytes) -> None:
    with ZipFile(BytesIO(content)) as archive:
        names = {name.casefold() for name in archive.namelist()}
        forbidden_prefixes = (
            "xl/externallinks/",
            "xl/querytables/",
            "xl/activex/",
            "xl/drawings/",
            "xl/charts/",
            "xl/tables/",
            "xl/pivottables/",
            "xl/pivotcache/",
            "xl/slicers/",
            "xl/comments",
        )
        assert not any(name.endswith(".bin") for name in names)
        assert "xl/connections.xml" not in names
        assert not any(name.startswith(forbidden_prefixes) for name in names)
        for name in archive.namelist():
            if name.endswith(".rels"):
                relationships = archive.read(name).lower()
                assert b'targetmode="external"' not in relationships
                assert b'/hyperlink"' not in relationships


def template_layout(worksheet) -> dict[str, object]:
    return {
        "merges": sorted(str(value) for value in worksheet.merged_cells.ranges),
        "widths": {column: worksheet.column_dimensions[column].width for column in "ABCDEFGHI"},
        "heights": {row: worksheet.row_dimensions[row].height for row in range(1, 57)},
        "print_area": worksheet.print_area,
        "orientation": worksheet.page_setup.orientation,
        "fit_width": worksheet.page_setup.fitToWidth,
        "fit_height": worksheet.page_setup.fitToHeight,
        "validations": [
            (str(rule.sqref), rule.type, rule.formula1)
            for rule in worksheet.data_validations.dataValidation
        ],
    }


def test_template_matches_official_form_format() -> None:
    workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert {column: worksheet.column_dimensions[column].width for column in "ABCDEFGHI"} == {
        "A": 2.6640625,
        "B": 10.83203125,
        "C": 8.83203125,
        "D": 7.1640625,
        "E": 17.83203125,
        "F": 6.1640625,
        "G": 12.5,
        "H": 7.33203125,
        "I": 9.5,
    }
    assert worksheet.row_dimensions[1].height == pytest.approx(28)
    assert worksheet.row_dimensions[2].height == pytest.approx(35)
    assert all(worksheet.row_dimensions[row].height == pytest.approx(30) for row in range(4, 55))
    assert worksheet["B1"].font.name == "微软雅黑"
    assert worksheet["B1"].font.bold is True
    assert worksheet["B55"].fill.fgColor.rgb == "FFB5C6EA"
    assert worksheet["G55"].number_format.startswith("_ \\¥*")
    workbook.close()


def test_generated_workbook_preserves_cells_order_totals_layout_and_safety(client_factory) -> None:
    before_hash = hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
    assert before_hash == CANONICAL_TEMPLATE_SHA256
    template_workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    template_worksheet = template_workbook[EXCEL_TEMPLATE.sheet_name]
    apply_output_page_setup(template_worksheet)
    expected_layout = template_layout(template_worksheet)
    template_workbook.close()

    result = generate(
        client_factory(),
        body(
            items=[
                item(
                    category="local_transport",
                    date="2026-07-06",
                    display_date="7月1日、7月6日",
                    description="酒店-项目-酒店",
                    amount="44.89",
                    receipt_count=4,
                ),
                item(
                    category="rail_fare",
                    date="2026-06-30",
                    description="北京南-合肥南",
                    amount="454.00",
                ),
                item(
                    category="rail_fare",
                    date="2026-07-07",
                    description="合肥南-北京南",
                    amount="473.50",
                ),
            ]
        ),
    )
    assert_archive_has_no_active_content(result.content)
    workbook = open_result(result)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert [worksheet[f"B{row}"].value for row in range(4, 8)] == [
        "火车票",
        "市内交通费",
        "火车票",
        "出差补助",
    ]
    assert worksheet["D7"].value == "6/30—7/7，共8天出差补助"
    assert [str(worksheet[f"G{row}"].value) for row in range(4, 8)] == [
        "454",
        "44.89",
        "473.5",
        "800",
    ]
    assert worksheet["C55"].value == "壹仟柒佰柒拾贰元叁角玖分"
    assert str(worksheet["G55"].value) == "1772.39"
    assert worksheet["I55"].value == 6
    assert template_layout(worksheet) == expected_layout
    assert not any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row)
    assert not any(cell.hyperlink or cell.comment for row in worksheet.iter_rows() for cell in row)
    assert not worksheet.tables and not worksheet._charts and not worksheet._images
    assert len(workbook._external_links) == 0
    workbook.close()
    assert hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest() == before_hash


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_formula_injection_is_saved_as_text(client_factory, prefix: str) -> None:
    payload = body(items=[item(display_date=f"  {prefix}DATE()", description=f"{prefix}DESC()")])
    payload["project"] = {"mode": "manual", "text": f"{prefix}PROJECT()"}
    workbook = open_result(
        generate(
            client_factory(),
            payload,
            employee_name=f"{prefix}NAME()",
            department_name=f"{prefix}DEPT()",
        )
    )
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    for coordinate in ("C2", "E2", "G2", "C4", "D4"):
        assert worksheet[coordinate].data_type == "s"
        assert worksheet[coordinate].value.startswith("'")
    assert not any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row)
    workbook.close()


def test_workbook_dynamically_extends_detail_rows(client_factory) -> None:
    result = generate(
        client_factory(),
        body(items=[item(description=f"费用 {index}") for index in range(75)]),
    )
    workbook = open_result(result)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert worksheet["D78"].value == "费用 74"
    assert worksheet["B79"].value == "出差补助"
    assert worksheet["B80"].value == "总合计"
    assert str(worksheet["G80"].value) == "875"
    assert worksheet["I80"].value == 75
    assert str(worksheet.data_validations.dataValidation[0].sqref) == "B4:B79"
    assert worksheet.print_area == "'费用报销模板'!$B$1:$I$80"
    workbook.close()


@pytest.mark.parametrize("damage", ["missing_sheet", "altered_merge", "print_titles"])
def test_invalid_template_returns_stable_error(client_factory, tmp_path: Path, damage: str) -> None:
    bad_template = tmp_path / f"{damage}.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    if damage == "missing_sheet":
        worksheet.title = "WrongSheet"
    elif damage == "altered_merge":
        worksheet.unmerge_cells("D4:F4")
    else:
        worksheet.print_title_rows = "1:4"
    workbook.save(bad_template)
    workbook.close()
    assert_template_invalid(client_factory, bad_template)


@pytest.mark.parametrize(
    "carrier",
    ["formula", "hyperlink", "defined_name", "conditional_formatting", "table", "chart", "comment"],
)
def test_template_active_content_is_rejected(client_factory, tmp_path: Path, carrier: str) -> None:
    bad_template = tmp_path / f"{carrier}.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    if carrier == "formula":
        worksheet["G4"] = "=1+1"
    elif carrier == "hyperlink":
        worksheet["D4"].hyperlink = "https://example.invalid/"
    elif carrier == "defined_name":
        workbook.defined_names.add(DefinedName("UnexpectedName", attr_text="'费用报销模板'!$B$2"))
    elif carrier == "conditional_formatting":
        worksheet.conditional_formatting.add(
            "G4", FormulaRule(formula=['WEBSERVICE("https://example.invalid")'])
        )
    elif carrier == "table":
        worksheet["A26"], worksheet["B26"] = "列一", "列二"
        worksheet["A27"], worksheet["B27"] = "值一", "值二"
        worksheet.add_table(Table(displayName="UnexpectedTable", ref="A26:B27"))
    elif carrier == "chart":
        chart = BarChart()
        chart.add_data(Reference(worksheet, min_col=7, min_row=3, max_row=4), titles_from_data=True)
        worksheet.add_chart(chart, "K2")
    else:
        worksheet["D4"].comment = Comment("不应存在的批注", "attacker")
    workbook.save(bad_template)
    workbook.close()
    assert_template_invalid(client_factory, bad_template)


@pytest.mark.parametrize(
    "member_name",
    [
        "xl/connections.xml",
        "xl/externalLinks/externalLink1.xml",
        "xl/queryTables/queryTable1.xml",
        "xl/vbaProject.bin",
    ],
)
def test_injected_active_ooxml_part_is_rejected(
    client_factory, tmp_path: Path, member_name: str
) -> None:
    bad_template = tmp_path / member_name.replace("/", "-")
    bad_template.write_bytes(TEMPLATE_PATH.read_bytes())
    with ZipFile(bad_template, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr(member_name, b"active-content")
    assert_template_invalid(client_factory, bad_template)


def test_filename_is_sanitized_and_rfc5987_encoded() -> None:
    filename = build_download_filename('测/试:*?"<>|用户', "预算/项目:一")
    assert filename.startswith("差旅费报销单-") and filename.endswith(".xlsx")
    assert not any(character in filename for character in '/\\:*?"<>|')
    header = content_disposition(filename)
    assert "filename=expense-report.xlsx" in header
    assert unquote(header.split("filename*=UTF-8''", 1)[1]) == filename
