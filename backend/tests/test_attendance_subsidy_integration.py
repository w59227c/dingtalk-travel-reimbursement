"""Calendar guards through real draft/Excel/submit APIs and a synthetic OA boundary."""

import asyncio
import json
from dataclasses import replace
from datetime import date
from io import BytesIO

import pytest
import test_oa_reimbursement_integration as integration
from conftest import mock_login
from openpyxl import load_workbook
from sqlalchemy import select
from test_attendance_subsidy import install_calendar

from app.integrations.dingtalk.workflow import FormOption
from app.models.reimbursement import ReimbursementDraft, ReimbursementSubmission
from app.schemas.expenses import TripPurpose
from app.services.attendance_calendar import delete_calendar, list_calendars
from app.services.oa_reimbursement_payload import ReimbursementSnapshot, build_snapshot


@pytest.fixture
def city_flow(client_factory, monkeypatch, request):
    async def manual_worker(_self, stop):
        await stop.wait()

    monkeypatch.setattr(integration.DurableOAReimbursementWorker, "run", manual_worker)
    client = client_factory(auth_mock_enabled=True, dingtalk_oa_worker_enabled=True)
    headers = {
        "X-CSRF-Token": mock_login(client)["csrfToken"],
        "Idempotency-Key": "86868686-8686-4868-8868-868686868686",
    }
    install_calendar(client)
    reimbursement, travel, _ = integration._schemas()
    city = FormOption(
        "境内市内项目出差（长/短期）", "境内市内项目出差（长/短期）", "synthetic-city"
    )
    reimbursement = integration._schema(
        reimbursement.process_code,
        tuple(
            replace(component, options=(city,))
            if component.component_id == "travel-type-id"
            else component
            for component in reimbursement.components
        ),
    )
    original = integration._draft_input
    overlapping = getattr(request, "param", False)
    instance_ids = (integration.TRAVEL_INSTANCE_ID,)
    if overlapping:
        instance_ids += (integration.TRAVEL_INSTANCE_ID_2,)

    # Only change the synthetic input fixture; all real validation remains active.
    def city_input(**kwargs):
        value = original(**kwargs)
        first = value.trip.model_copy(update={
            "trip_type": TripPurpose.SAME_CITY_PROJECT,
            "attendance_dates": [date(2026, 9, 1), date(2026, 9, 3)],
        })
        if overlapping:
            return value.model_copy(update={"trip": None, "trips": [
                first.model_copy(update={"related_approval_id": instance_ids[0]}),
                first.model_copy(update={
                    "related_approval_id": instance_ids[1],
                    "attendance_dates": [date(2026, 9, 2), date(2026, 9, 3)],
                }),
            ]})
        return value.model_copy(
            update={"trip": first}
        )

    monkeypatch.setattr(integration, "_draft_input", city_input)
    workflow = integration.LocalWorkflowBoundary(
        reimbursement, travel, membership_pages=(instance_ids,)
    )
    draft_id, _ = integration._persist_ready_draft(
        client,
        workflow,
        city,
        subsidy_only=True,
        with_materials=False,
        travel_instance_ids=instance_ids,
    )
    return client, headers, draft_id, workflow


@pytest.mark.parametrize("city_flow", [True], indirect=True)
def test_overlapping_approvals_preview_and_submission_use_one_exact_date_union(city_flow):
    client, headers, draft_id, workflow = city_flow
    preview = client.post(
        f"/api/reimbursements/drafts/{draft_id}/excel-preview",
        json={"expectedRevision": 4}, headers=headers,
    )
    assert preview.status_code == 200, preview.text
    workbook = load_workbook(BytesIO(preview.content))
    try:
        text = "\n".join(str(cell.value) for row in workbook.active for cell in row if cell.value)
        assert text.count("出勤日期（2026年9月）：1—3日") == 1
        assert "出勤3天 × 50.00元/天 = 150.00元" in text
    finally:
        workbook.close()
    response = client.post(
        f"/api/oa/reimbursements/{draft_id}/submit",
        json={"expectedRevision": 4}, headers=headers,
    )
    assert response.status_code == 202, response.text
    identifier = response.json()["data"]["submissionId"]
    with client.app.state.database_session_factory() as database:
        value = json.loads(database.get(ReimbursementSubmission, identifier).form_snapshot_json)
        assert len(value["input"]["trips"]) == 2
        assert len(value["subsidies"]) == 1
        assert value["totals"]["subsidyTotal"] == "150.00"
        description = next(
            field["value"] for field in value["formValues"] if field["logicalKey"] == "description"
        )
        assert description.count("出勤日期（2026年9月）：1—3日") == 1
    worker = integration._worker(client, workflow, integration.LocalStorageBoundary())
    assert asyncio.run(worker.run_once())
    assert client.get(
        f"/api/oa/reimbursements/submissions/{identifier}"
    ).json()["data"]["status"] == "SUBMITTED"


def remove_calendar(client):
    with client.app.state.database_session_factory() as database:
        calendar = next(item for item in list_calendars(database) if item["year"] == 2026)
        delete_calendar(database, 2026, calendar["revision"])


def test_missing_calendar_blocks_excel_review_and_submit_without_locking_or_losing_input(city_flow):
    client, headers, draft_id, workflow = city_flow
    remove_calendar(client)
    for path in (
        f"/api/reimbursements/drafts/{draft_id}/excel-preview",
        f"/api/reimbursements/drafts/{draft_id}/review",
        f"/api/oa/reimbursements/{draft_id}/submit",
    ):
        response = client.post(path, json={"expectedRevision": 4}, headers=headers)
        assert response.status_code == 422, response.text
        assert response.json()["error"]["code"] == "ATTENDANCE_CALENDAR_REQUIRED"
    restored = client.get(f"/api/reimbursements/drafts/{draft_id}").json()["data"]
    assert restored["input"]["trip"]["attendanceDates"] == ["2026-09-01", "2026-09-03"]
    with client.app.state.database_session_factory() as database:
        draft = database.get(ReimbursementDraft, draft_id)
        assert draft.locked_at is None and draft.revision == 4
        assert database.scalar(select(ReimbursementSubmission.id)) is None
    assert workflow.create_calls == 0
    install_calendar(client)
    response = client.post(
        f"/api/reimbursements/drafts/{draft_id}/excel-preview",
        json={"expectedRevision": 4},
        headers=headers,
    )
    assert response.status_code == 200, response.text


@pytest.mark.parametrize("legacy", [False, True])
def test_locked_city_snapshot_preview_idempotency_and_worker_ignore_later_calendar_removal(
    city_flow, monkeypatch, legacy,
):
    client, headers, draft_id, workflow = city_flow
    path = f"/api/oa/reimbursements/{draft_id}/submit"
    # Model an old app's completed lock, then resume with all new code restored.
    def legacy_build(source):
        value = build_snapshot(source).model_dump(mode="json", by_alias=True)
        value.pop("attendanceDescriptionFormat", None)
        for field in value["formValues"]:
            if field["logicalKey"] == "description":
                field["value"] = (
                    "市内出差补助：出勤2天 × 50.00元/天 = 100.00元\n"
                    "出勤日期：2026-09-01、2026-09-03\n合计：100.00"
                )
        return ReimbursementSnapshot.model_validate(value)

    with monkeypatch.context() as patch:
        if legacy:
            patch.setattr("app.api.reimbursement_submissions.build_snapshot", legacy_build)
        submitted = client.post(path, json={"expectedRevision": 4}, headers=headers)
    assert submitted.status_code == 202, submitted.text
    submission_id = submitted.json()["data"]["submissionId"]
    with client.app.state.database_session_factory() as database:
        row = database.get(ReimbursementSubmission, submission_id)
        frozen_json, frozen_digest = row.form_snapshot_json, row.snapshot_sha256
    expected = (
        "出勤日期：2026-09-01、2026-09-03" if legacy else "出勤日期（2026年9月）：1日、3日"
    )
    assert expected in frozen_json
    remove_calendar(client)
    repeated = client.post(path, json={"expectedRevision": 4}, headers=headers)
    assert repeated.status_code == 202, repeated.text
    assert repeated.json()["data"]["submissionId"] == submission_id
    locked = client.get(f"/api/reimbursements/drafts/{draft_id}").json()["data"]
    assert locked["totals"]["subsidyTotal"] == "100.00"
    preview = client.post(
        f"/api/reimbursements/drafts/{draft_id}/excel-preview",
        json={"expectedRevision": locked["revision"]},
        headers=headers,
    )
    assert preview.status_code == 200, preview.text

    def assert_workbook(content):
        workbook = load_workbook(BytesIO(content))
        try:
            texts = [str(cell.value) for row in workbook.active for cell in row if cell.value]
            assert any(expected in text for text in texts)
            if legacy:
                assert not any("出勤日期（" in text for text in texts)
        finally:
            workbook.close()

    assert_workbook(preview.content)
    storage = integration.LocalStorageBoundary()
    worker = integration._worker(client, workflow, storage)
    assert asyncio.run(worker.run_once())
    status = client.get(f"/api/oa/reimbursements/submissions/{submission_id}").json()["data"]
    assert status["status"] == "SUBMITTED", status
    assert workflow.create_calls == 1
    assert_workbook(storage.put_files[0].content)
    frozen_description = next(
        field["value"] for field in json.loads(frozen_json)["formValues"]
        if field["logicalKey"] == "description"
    )
    assert any(field.value == frozen_description for field in workflow.created_command.form_values)
    after = client.post(
        f"/api/reimbursements/drafts/{draft_id}/excel-preview",
        json={"expectedRevision": locked["revision"]}, headers=headers,
    )
    assert after.status_code == 200, after.text
    assert_workbook(after.content)
    repeated = client.post(path, json={"expectedRevision": 4}, headers=headers)
    assert repeated.json()["data"]["submissionId"] == submission_id
    assert not asyncio.run(worker.run_once())
    assert workflow.create_calls == 1
    with client.app.state.database_session_factory() as database:
        row = database.get(ReimbursementSubmission, submission_id)
        assert (row.form_snapshot_json, row.snapshot_sha256) == (frozen_json, frozen_digest)
