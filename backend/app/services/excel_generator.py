from __future__ import annotations

import logging
import math
import re
import unicodedata
from copy import copy
from dataclasses import dataclass, replace
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import ApiError
from app.domain.categories import CATEGORY_BY_ID, ExpenseCategory
from app.domain.expenses import ExpenseTotals
from app.domain.money import quantize_money
from app.domain.subsidy import SubsidyCalculation
from app.excel.template_contract import (
    EXCEL_TEMPLATE,
    ExcelSheetLayout,
    apply_output_page_setup,
    output_layout_for_line_count,
    template_invalid,
    validate_workbook_safety,
    validate_xlsx_archive,
)
from app.schemas.excel import ExcelExpenseItemInput
from app.schemas.expenses import TripInput
from app.services.subsidy_description import subsidy_explanation

logger = logging.getLogger(__name__)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TOTAL_AMOUNT_NUMBER_FORMAT = '_ \\¥* #,##0.00_ ;_ \\¥* \\-#,##0.00_ ;_ \\¥* "-"??_ ;_ @_ '
_ILLEGAL_XML_CONTROLS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_INVALID_FILENAME_CHARACTERS = re.compile(r"[\\/:*?\"<>|\x00-\x1f\x7f]")
_DANGEROUS_FORMULA_PREFIXES = ("=", "+", "-", "@")
_MAX_ROW_HEIGHT = 409.5
_MAX_CELL_TEXT_LENGTH = 32_767


@dataclass(frozen=True, slots=True)
class ResolvedProject:
    display_text: str
    filename_component: str


@dataclass(frozen=True, slots=True)
class WorkbookResult:
    content: bytes
    filename: str


@dataclass(frozen=True, slots=True)
class OutputLine:
    category: ExpenseCategory
    sort_date: date
    display_date: str
    description: str
    amount: Decimal
    receipt_count: int
    subsidy: bool = False
    continuation: bool = False
    detail_lines: tuple[str, ...] = ()


def sanitize_external_text(value: str, *, max_length: int) -> str:
    cleaned = _ILLEGAL_XML_CONTROLS.sub(" ", str(value))[:max_length]
    if cleaned.lstrip().startswith(_DANGEROUS_FORMULA_PREFIXES):
        cleaned = "'" + cleaned
    return cleaned


def write_safe_text(cell: Cell, value: str, *, max_length: int) -> None:
    cell.value = sanitize_external_text(value, max_length=max_length)
    cell.data_type = "s"


def sanitize_filename_component(value: str, *, max_length: int = 60) -> str:
    cleaned = _INVALID_FILENAME_CHARACTERS.sub("_", str(value)).strip(" .")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned[:max_length].rstrip(" .") or "未命名"


def build_download_filename(employee_name: str, project_component: str) -> str:
    employee = sanitize_filename_component(employee_name)
    project = sanitize_filename_component(project_component)
    return f"差旅费报销单-{employee}-{project}.xlsx"


def content_disposition(filename: str) -> str:
    encoded = quote(filename, safe="")
    return f"attachment; filename=expense-report.xlsx; filename*=UTF-8''{encoded}"


def load_validated_template(template_path: Path) -> tuple[Workbook, Worksheet]:
    """Load the canonical workbook only after the complete fail-closed contract check."""

    validate_xlsx_archive(template_path)
    try:
        workbook = load_workbook(template_path, read_only=False, data_only=False, keep_links=True)
    except Exception as exc:
        logger.error(
            "Failed to load Excel template",
            extra={"exception_type": type(exc).__name__},
        )
        raise template_invalid() from None

    try:
        if workbook.sheetnames == [EXCEL_TEMPLATE.sheet_name]:
            apply_output_page_setup(workbook[EXCEL_TEMPLATE.sheet_name])
        worksheet = validate_workbook_safety(workbook)
        return workbook, worksheet
    except Exception:
        workbook.close()
        raise


def _compact_date(value: date) -> str:
    return f"{value.month}月{value.day}日"


def _compact_date_range(start: date, end: date) -> str:
    if start.year == end.year:
        return f"{start.month}/{start.day}—{end.month}/{end.day}"
    return f"{start.year}/{start.month}/{start.day}—{end.year}/{end.month}/{end.day}"


def _compact_days(value: Decimal) -> str:
    if value == value.to_integral_value():
        return str(int(value))
    return format(value.normalize(), "f")


def _display_date(value: date, requested: str) -> str:
    return _compact_date(value) if requested.strip() == value.isoformat() else requested


def _subsidy_description(trip: TripInput, subsidy: SubsidyCalculation) -> str:
    period = _compact_date_range(trip.start_date, trip.end_date)
    exclusion = (
        f"（扣除{_compact_days(subsidy.excluded_days)}天非补助时段）"
        if subsidy.excluded_days is not None
        else ""
    )
    return f"{period}，共{_compact_days(subsidy.effective_days)}天出差补助{exclusion}"


def _ordered_lines(
    items: list[ExcelExpenseItemInput],
    trip: TripInput | None,
    subsidy: SubsidyCalculation | None,
    trips: list[TripInput] | tuple[TripInput, ...] = (),
    subsidies: list[SubsidyCalculation] | tuple[SubsidyCalculation, ...] = (),
    detailed_subsidies: bool = True,
) -> list[OutputLine]:
    indexed_lines = [
        (
            item.date,
            index,
            OutputLine(
                category=item.category,
                sort_date=item.date,
                display_date=_display_date(item.date, item.display_date),
                description=item.description,
                amount=item.amount,
                receipt_count=item.receipt_count,
            ),
        )
        for index, item in enumerate(items)
    ]
    if trips and subsidies:
        subsidy_pairs = list(zip(trips, subsidies, strict=True))
    elif trip is not None and subsidy is not None:
        subsidy_pairs = [(trip, subsidy)]
    else:
        # An overseas approval is a valid reimbursement relation but creates
        # no subsidy line.
        subsidy_pairs = []
    for index, (subsidy_trip, subsidy_item) in enumerate(subsidy_pairs):
        explanation = (
            subsidy_explanation(subsidy_trip, subsidy_item)
            if detailed_subsidies
            and (subsidy_trip.excluded_periods or subsidy_trip.attendance_dates is not None)
            else ()
        )
        indexed_lines.append(
            (
                subsidy_trip.end_date,
                len(items) + index,
                OutputLine(
                    category=ExpenseCategory.SUBSIDY,
                    sort_date=subsidy_trip.end_date,
                    display_date=_compact_date(subsidy_trip.end_date),
                    description=explanation[0] if explanation else _subsidy_description(
                        subsidy_trip, subsidy_item
                    ),
                    amount=subsidy_item.total,
                    receipt_count=0,
                    subsidy=True,
                    detail_lines=explanation[1:],
                ),
            )
        )
    return [line for _, _, line in sorted(indexed_lines, key=lambda item: (item[0], item[1]))]


def _description_height(text: str, width: float) -> float:
    wrapped_lines = sum(
        max(1, math.ceil(sum(
            2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
            for char in paragraph
        ) / max(1, width - 2)))
        for paragraph in text.split("\n")
    )
    return max(30, wrapped_lines * 18 + 6)


def _expand_subsidy_descriptions(lines: list[OutputLine], width: float) -> list[OutputLine]:
    """Keep an adjustment summary in one cell; continue only past Excel's height limit."""
    output = []
    for line in lines:
        if not line.detail_lines:
            output.append(line)
            continue
        chunks: list[str] = []
        current = line.description
        for detail in line.detail_lines:
            combined = f"{current}\n{detail}"
            if _description_height(combined, width) > _MAX_ROW_HEIGHT:
                chunks.append(current)
                current = detail
            else:
                current = combined
        chunks.append(current)
        output.append(replace(line, description=chunks[0]))
        output.extend(
            replace(
                line, description=chunk, display_date="", amount=Decimal(0), receipt_count=0,
                continuation=True, detail_lines=(),
            )
            for chunk in chunks[1:]
        )
    return output


def _copy_row(
    worksheet: Worksheet,
    *,
    source_row: int,
    target_row: int,
    copy_values: bool,
) -> None:
    """Copy the maintained row's presentation without introducing formulas."""

    for column in range(1, worksheet.max_column + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
        target.value = source.value if copy_values else None

    source_dimension = worksheet.row_dimensions[source_row]
    target_dimension = worksheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden if copy_values else False
    target_dimension.outlineLevel = source_dimension.outlineLevel
    target_dimension.collapsed = source_dimension.collapsed
    target_dimension.thickTop = source_dimension.thickTop
    target_dimension.thickBot = source_dimension.thickBot


def _prepare_output_layout(
    worksheet: Worksheet,
    line_count: int,
) -> ExcelSheetLayout:
    layout = output_layout_for_line_count(line_count)
    if layout.detail_end_row > EXCEL_TEMPLATE.detail_end_row:
        # Preserve the template total row before reusing its original position
        # as the first dynamically appended detail row.
        _copy_row(
            worksheet,
            source_row=EXCEL_TEMPLATE.total_row,
            target_row=layout.total_row,
            copy_values=True,
        )
        worksheet.merge_cells(
            start_row=layout.total_row,
            start_column=3,
            end_row=layout.total_row,
            end_column=6,
        )
        worksheet.unmerge_cells(
            start_row=EXCEL_TEMPLATE.total_row,
            start_column=3,
            end_row=EXCEL_TEMPLATE.total_row,
            end_column=6,
        )

        for row in range(EXCEL_TEMPLATE.total_row, layout.detail_end_row + 1):
            _copy_row(
                worksheet,
                source_row=EXCEL_TEMPLATE.detail_end_row,
                target_row=row,
                copy_values=False,
            )
            worksheet.merge_cells(
                start_row=row,
                start_column=4,
                end_row=row,
                end_column=6,
            )

    validation = worksheet.data_validations.dataValidation[0]
    validation.sqref = f"B{EXCEL_TEMPLATE.detail_start_row}:B{layout.detail_end_row}"
    apply_output_page_setup(worksheet, layout)
    return layout


def generate_expense_workbook(
    *,
    template_path: Path,
    employee_name: str,
    department_name: str,
    project: ResolvedProject,
    trip: TripInput | None,
    items: list[ExcelExpenseItemInput],
    subsidy: SubsidyCalculation | None,
    totals: ExpenseTotals,
    trips: list[TripInput] | tuple[TripInput, ...] = (),
    subsidies: list[SubsidyCalculation] | tuple[SubsidyCalculation, ...] = (),
    detailed_subsidies: bool = True,
) -> WorkbookResult:
    workbook, worksheet = load_validated_template(template_path)
    try:
        write_safe_text(
            worksheet[EXCEL_TEMPLATE.employee_name_cell],
            employee_name,
            max_length=128,
        )
        write_safe_text(
            worksheet[EXCEL_TEMPLATE.department_cell],
            department_name,
            max_length=255,
        )
        write_safe_text(
            worksheet[EXCEL_TEMPLATE.project_cell],
            project.display_text,
            max_length=2048,
        )
        # The budget label may be longer than a short project name. Keep the
        # template's wrapping and expand its header row for readable printouts.
        project_cell = worksheet[EXCEL_TEMPLATE.project_cell]
        text_width = sum(
            2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
            for char in project.display_text
        )
        column_width = sum(worksheet.column_dimensions[key].width for key in ("G", "H", "I"))
        header_height = math.ceil(text_width / max(1, column_width - 2)) * 16
        worksheet.row_dimensions[project_cell.row].height = min(
            409.5, max(worksheet.row_dimensions[project_cell.row].height or 35, header_height)
        )

        description_width = sum(
            worksheet.column_dimensions[key].width for key in ("D", "E", "F")
        )
        lines = _expand_subsidy_descriptions(
            _ordered_lines(items, trip, subsidy, trips, subsidies, detailed_subsidies),
            description_width,
        )
        layout = _prepare_output_layout(worksheet, len(lines))
        for offset, line in enumerate(lines):
            row = EXCEL_TEMPLATE.detail_start_row + offset
            if line.subsidy and detailed_subsidies and (line.continuation or line.detail_lines):
                description_cell = worksheet[f"{EXCEL_TEMPLATE.description_column}{row}"]
                alignment = copy(description_cell.alignment)
                alignment.wrap_text = True
                alignment.horizontal = "left"
                alignment.vertical = "top"
                description_cell.alignment = alignment
                worksheet.row_dimensions[row].height = min(
                    _MAX_ROW_HEIGHT, _description_height(line.description, description_width)
                )
            if line.continuation:
                for column in (
                    EXCEL_TEMPLATE.category_column, EXCEL_TEMPLATE.date_column,
                    EXCEL_TEMPLATE.amount_column, EXCEL_TEMPLATE.currency_column,
                    EXCEL_TEMPLATE.receipt_count_column,
                ):
                    worksheet[f"{column}{row}"] = None
                write_safe_text(
                    worksheet[f"{EXCEL_TEMPLATE.description_column}{row}"],
                    line.description, max_length=_MAX_CELL_TEXT_LENGTH,
                )
                worksheet.row_dimensions[row].hidden = False
                continue
            category_name = CATEGORY_BY_ID[line.category].name
            write_safe_text(
                worksheet[f"{EXCEL_TEMPLATE.category_column}{row}"],
                category_name,
                max_length=20,
            )
            write_safe_text(
                worksheet[f"{EXCEL_TEMPLATE.date_column}{row}"],
                line.display_date,
                max_length=100,
            )
            write_safe_text(
                worksheet[f"{EXCEL_TEMPLATE.description_column}{row}"],
                line.description,
                max_length=_MAX_CELL_TEXT_LENGTH if line.detail_lines else 500,
            )
            worksheet[f"{EXCEL_TEMPLATE.amount_column}{row}"] = quantize_money(line.amount)
            write_safe_text(
                worksheet[f"{EXCEL_TEMPLATE.currency_column}{row}"],
                "人民币",
                max_length=10,
            )
            worksheet[f"{EXCEL_TEMPLATE.receipt_count_column}{row}"] = line.receipt_count
            worksheet.row_dimensions[row].hidden = False

        first_unused = EXCEL_TEMPLATE.detail_start_row + len(lines)
        for row in range(first_unused, layout.detail_end_row + 1):
            worksheet.row_dimensions[row].hidden = True

        write_safe_text(
            worksheet[layout.uppercase_amount_cell],
            totals.uppercase_amount,
            max_length=100,
        )
        total_amount_cell = worksheet[layout.total_amount_cell]
        total_amount_cell.value = quantize_money(totals.total_amount)
        total_amount_cell.number_format = TOTAL_AMOUNT_NUMBER_FORMAT
        worksheet[layout.total_receipt_count_cell] = totals.receipt_count

        validate_workbook_safety(workbook, layout)
        output = BytesIO()
        workbook.save(output)
        validate_xlsx_archive(output.getvalue())
        return WorkbookResult(
            content=output.getvalue(),
            filename=build_download_filename(employee_name, project.filename_component),
        )
    except ApiError:
        raise
    except Exception as exc:
        logger.error(
            "Failed to generate Excel workbook",
            extra={"exception_type": type(exc).__name__},
        )
        raise ApiError(
            "EXCEL_GENERATION_FAILED",
            "Excel 生成失败，请稍后重试",
            500,
        ) from None
    finally:
        workbook.close()
